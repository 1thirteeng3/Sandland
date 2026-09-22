from __future__ import annotations
import csv
import hashlib
import importlib
import json
from collections import Counter, defaultdict, deque
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

from plan_base import ROOT, SPRINTS, RELEASES, REQUIREMENTS, ADRS
from guide import OPENING, CLOSING, ENGINES, RISKS

for n in range(1,7): importlib.import_module(f'phase_{n:02}')
SPRINTS.sort(key=lambda s:s['n'])
OUT=ROOT/'planejamento'
OUT.mkdir(exist_ok=True)
MD=ROOT/'Sandland-Roadmap-v0-a-v2.md'
XLSX=ROOT/'Sandland-Backlog.xlsx'
CSV=ROOT/'Sandland-Backlog.csv'
STATES=['Não iniciado','Pronto','Em andamento','Em revisão','Bloqueado','Concluído','Dispensado (ADR)']
TRACK=['status','reviewer','accepted_at','evidence_obtained','adr_justification','notes','owner']

# Preserva acompanhamento por ID; o conteúdo técnico é regenerado dos dados versionados.
prior={}
prior_adrs={}
if XLSX.exists():
    old=load_workbook(XLSX,read_only=True,data_only=False)
    if 'Tarefas' in old.sheetnames:
        for row in old['Tarefas'].iter_rows(min_row=2,values_only=True):
            if len(row)>=25 and row[1]:
                prior[str(row[1])]=dict(status=row[13],owner=row[17],reviewer=row[18],accepted_at=row[19],evidence_obtained=row[21],adr_justification=row[22],notes=row[24])
    if 'ADRs' in old.sheetnames:
        for row in old['ADRs'].iter_rows(min_row=2,values_only=True):
            if len(row)>=7 and row[0]:
                prior_adrs[str(row[0])]=[row[4] or 'A decidir',row[5] or 'Autor',row[6] or '']
    old.close()

all_tasks=[]
for s in SPRINTS:
    implementation_ids=[]
    for i,t in enumerate(s['tasks'],1):
        tid=f"SL-{s['id']}-{i:02d}"
        deps=[implementation_ids[-1]] if implementation_ids else [f'SL-{d}-GATE' for d in s['deps']]
        implementation_ids.append(tid)
        r=dict(t)
        r.update(id=tid,seq=len(all_tasks)+1,sprint=s['id'],release=s['release'],component=s['component'],nature=s['nature'],requirements=s['requirements'],dependencies=deps,risk=s['risk'],priority='P0' if s['risk']=='Crítico' or t['kind'] in ['Decisão','Contrato'] else 'P2' if s['n']>=64 else 'P1',applicability='Condicional à ratificação do escopo pós-v1 (ADR-017/018)' if s['n']>=64 else 'Obrigatória no escopo da versão',owner='Autor' if t['kind'] in ['Decisão','Governança','Licença','Release'] else 'Autor + agente de IA',status='Não iniciado',reviewer='',accepted_at='',evidence_expected=f"reports/{s['id']}/{tid}.md",evidence_obtained='',adr_justification='',notes='')
        for k,v in prior.get(tid,{}).items():
            if v is not None: r[k]=v
        if r['status'] not in STATES: raise ValueError(f'Status inválido: {tid}: {r["status"]}')
        all_tasks.append(r)
    gateid=f"SL-{s['id']}-GATE"
    gate=dict(id=gateid,seq=len(all_tasks)+1,sprint=s['id'],release=s['release'],component=s['component'],nature='Aceite humano',kind='Gate',title=f"Aceitar {s['id']} — {s['title']}",execution='Revisar as dez tarefas, executar a demonstração do sprint, confrontar critérios e registrar decisão go/no-go.',output=f"reports/{s['id']}/acceptance.md",acceptance=s['demo']+' Todas as dependências aceitas; evidências vinculadas a revisão/ambiente; nenhum blocker de integridade/segurança.',requirements=s['requirements'],dependencies=implementation_ids,risk=s['risk'],priority='P0',applicability='Obrigatória para fechar a versão aprovada',owner='Autor',status='Não iniciado',reviewer='',accepted_at='',evidence_expected=f"reports/{s['id']}/acceptance.md",evidence_obtained='',adr_justification='',notes='Uma conclusão negativa fundamentada pode fechar um spike, mas não aprova o motor rejeitado.' if s['nature']=='Spike de motor' else '')
    for k,v in prior.get(gateid,{}).items():
        if v is not None: gate[k]=v
    all_tasks.append(gate)
    s['task_ids']=implementation_ids
    s['gate_id']=gateid

byid={t['id']:t for t in all_tasks}
by_sprint={s['id']:s for s in SPRINTS}
reqmap={r[0]:r for r in REQUIREMENTS}
assert len(byid)==len(all_tasks)
assert [s['n'] for s in SPRINTS]==list(range(1,73))
assert all(len(s['tasks'])==10 for s in SPRINTS)
assert all(d in byid for t in all_tasks for d in t['dependencies'])
assert all(r in reqmap for t in all_tasks for r in t['requirements'])
assert all(r[0] in {x for s in SPRINTS for x in s['requirements']} for r in REQUIREMENTS)
# Validação acíclica do grafo completo, incluindo gates.
indeg={t['id']:len(set(t['dependencies'])) for t in all_tasks}
out=defaultdict(list)
for t in all_tasks:
    for d in set(t['dependencies']): out[d].append(t['id'])
q=deque(x for x,n in indeg.items() if n==0)
visited=[]
while q:
    x=q.popleft(); visited.append(x)
    for y in out[x]:
        indeg[y]-=1
        if indeg[y]==0:q.append(y)
assert len(visited)==len(all_tasks), 'Ciclo no grafo de dependências'
for tid in visited:
    t=byid[tid]
    has_date=isinstance(t['accepted_at'],(date,datetime)) or (isinstance(t['accepted_at'],(int,float)) and t['accepted_at']>0)
    has_evidence=bool(t['reviewer'] and has_date and t['evidence_obtained'])
    if t['status']=='Concluído':
        t['closure']='OK' if has_evidence and all(byid[d]['closure'] in ['OK','DISPENSADO'] for d in t['dependencies']) else 'REVISAR'
    elif t['status']=='Dispensado (ADR)':
        t['closure']='DISPENSADO' if has_evidence and t['applicability'].startswith('Condicional') and t['adr_justification'] else 'REVISAR'
    else:t['closure']='PENDENTE'


def fmt(x):
    if isinstance(x,(date,datetime)):return x.isoformat()
    return str(x or '')

def cell(x): return fmt(x).replace('|','\\|').replace('\n','<br>')

def mdtable(headers,rows):
    return '| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(cell(x) for x in r)+' |' for r in rows)

engine_rows=[(a,b,c,d,e,f) for a,b,c,d,e,f,u in ENGINES]
release_rows=[(r[0],f'S{r[1]:02d}–S{r[2]:02d}' if r[1]!=r[2] else f'S{r[1]:02d}',r[3],r[4],r[5]) for r in RELEASES]
sprintlinks='\n'.join(f"- **{r[0]}:** "+' · '.join(f"[{s['id']}](#{s['id'].lower()})" for s in SPRINTS if s['release']==r[0]) for r in RELEASES)
opening=OPENING
for key,value in {'sprints':len(SPRINTS),'tasks':sum(len(s['tasks']) for s in SPRINTS),'gates':len(SPRINTS),'versions':len(RELEASES),'items':len(all_tasks),'engines_table':mdtable(['Motor','Sprints','Integração','Reaproveitar','Excluir','Responsabilidade própria'],engine_rows),'releases_table':mdtable(['Versão','Sprints','Entrega','Maturidade','Resultado verificável'],release_rows),'sprint_links':sprintlinks}.items():
    opening=opening.replace('{{'+key+'}}',str(value))
parts=[opening]
for rel in RELEASES:
    parts.append(f"\n# {rel[0]} — {rel[3]}\n\n**Maturidade:** {rel[4]}  \n**Entrega acumulada:** {rel[5]}\n")
    for s in (x for x in SPRINTS if x['release']==rel[0]):
        deps='; '.join(f"{d} ({by_sprint[d]['title']})" for d in s['deps']) or 'Nenhum sprint anterior; requer acesso aos documentos de base.'
        parts.append(f"\n<a id=\"{s['id'].lower()}\"></a>\n\n## {s['id']} — {s['title']}\n\n- **Objetivo:** {s['objective']}\n- **Componente / forma:** {s['component']} / {s['nature']}.\n- **Pré-requisitos técnicos:** {deps}\n- **Entregável do sprint:** {s['deliverable']}\n- **Demonstração exigida:** {s['demo']}\n- **Requisitos:** {', '.join(s['requirements'])}. **Risco de integração:** {s['risk']}.\n")
        if s['note']:parts.append(f"\n> {s['note']}\n")
        for tid in s['task_ids']:
            t=byid[tid]
            parts.append(f"\n### {tid} — {t['title']}\n\n- **Executar:** {t['execution']}\n- **Produzir:** `{t['output']}`.\n- **Aceitar quando:** {t['acceptance']}\n- **Controle:** {t['kind']}; {t['priority']} na versão; depende de {', '.join(t['dependencies']) or 'Ready de S01'}; evidência em `{t['evidence_expected']}`.\n")
        parts.append(f"\n### {s['gate_id']} — Encerramento por aceite\n\nRevisão humana das dez tarefas e da demonstração: **{s['demo']}**\n\nRegistrar commit, ambiente, testes e pendências em `reports/{s['id']}/acceptance.md`. Só fechar quando os critérios particulares e a Definition of Done forem satisfeitos. Se houver blocker, manter o sprint aberto e registrar ação corretiva.\n\n---\n")

req_rows=[]
for rid,area,desc,source,accept in REQUIREMENTS:
    ss=[s['id'] for s in SPRINTS if rid in s['requirements']]
    req_rows.append((rid,area,desc,source,', '.join(ss),accept))
closing=CLOSING.replace('{{requirements_table}}',mdtable(['ID','Área','Requisito','Origem','Sprints','Verificação'],req_rows)).replace('{{adrs_table}}',mdtable(['ID','Decisão','Sprint','Regra'],ADRS)).replace('{{risks_table}}',mdtable(['ID','Risco','Severidade','Onde','Resposta'],RISKS))
parts.append(closing)
md=''.join(parts)
assert '{{' not in md
assert sum(line.startswith('```') for line in md.splitlines())%2==0
assert all(md.count('### '+t['id']+' —')==1 for t in all_tasks)
MD.write_text(md,encoding='utf-8')

source_files=[ROOT/'uploads/Documento de Arquitetura Técnica e Especificação de Produto.md',ROOT/'Sandland-RFC-Integracao-e-Desconstrucao.md']
source_manifest=[dict(file=str(p.relative_to(ROOT)),sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in source_files]
data=dict(plan_version='1.0',created='2026-09-22',capacity='Solo + agentes de IA',cadence='Por critérios de aceite; sem timebox',horizon='v2.0.0',sources=source_manifest,releases=RELEASES,sprints=SPRINTS,tasks=all_tasks,requirements=REQUIREMENTS,adrs=ADRS,adr_tracking=prior_adrs,engines=ENGINES,risks=RISKS)
(OUT/'Sandland-Roadmap-dados.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,default=fmt)+'\n',encoding='utf-8')

headers=['Ordem','ID','Versão','Sprint','Componente','Forma de integração','Tipo','Tarefa','Passo de execução','Artefato a produzir','Critério de aceite','Dependências','Requisitos','Status','Prioridade na versão','Risco','Aplicabilidade','Responsável','Revisor humano','Data de aceite','Evidência exigida','Evidência obtida','ADR / justificativa','Fechamento','Observações']

def taskrow(t):
    return [t['seq'],t['id'],t['release'],t['sprint'],t['component'],t['nature'],t['kind'],t['title'],t['execution'],t['output'],t['acceptance'],'; '.join(t['dependencies']),'; '.join(t['requirements']),t['status'],t['priority'],t['risk'],t['applicability'],t['owner'],t['reviewer'],t['accepted_at'],t['evidence_expected'],t['evidence_obtained'],t['adr_justification'],t['closure'],t['notes']]

with CSV.open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f)
    w.writerow(headers)
    for t in all_tasks:w.writerow(taskrow(t))

# Workbook: células de acompanhamento permanecem editáveis; fórmulas não certificam evidência.
wb=Workbook()
wb.remove(wb.active)
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
wb.properties.title='Sandland — Roadmap e Backlog até v2.0'
wb.properties.subject='72 sprints por aceite; implementação solo com agentes de IA'
wb.properties.creator='Planejamento Sandland'
wb.properties.description='Plano futuro: tarefas e gates não equivalem a implementação concluída.'
NAVY='102A43'; TEAL='0F766E'; BLUE='2563EB'; GRAY='52667A'; LIGHT='EAF3F8'; WHITE='FFFFFF'; RED='FDE8E8'; GREEN='DCFCE7'; AMBER='FEF3C7'
basefont=Font(name='Calibri',size=11,color=NAVY)
headfill=PatternFill('solid',fgColor=NAVY)
linkfont=Font(name='Calibri',size=11,color=BLUE,underline='single')


def tabular(name,cols,rows,widths,table_name=None):
    ws=wb.create_sheet(name)
    ws.append(cols)
    for row in rows:ws.append(list(row))
    ws.freeze_panes='B2'
    ws.sheet_view.showGridLines=False
    for c in ws[1]:
        c.fill=headfill;c.font=Font(name='Calibri',bold=True,color=WHITE,size=11)
        c.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=34
    for row in ws.iter_rows(min_row=2):
        for c in row:c.font=basefont;c.alignment=Alignment(vertical='top',wrap_text=True)
    for i,width in enumerate(widths,1):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(i)].width=width
    if rows and table_name:
        table=Table(displayName=table_name,ref=ws.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
        ws.add_table(table)
    elif rows:ws.auto_filter.ref=ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    return ws

readme=wb.create_sheet('LEIA-ME')
readme.sheet_view.showGridLines=False
readme.merge_cells('A1:H2');readme['A1']='SANDLAND / PLANO DE IMPLEMENTAÇÃO'
readme['A1'].fill=headfill;readme['A1'].font=Font(name='Calibri',size=22,bold=True,color=WHITE);readme['A1'].alignment=Alignment(vertical='center')
info=[
 ('Modelo de trabalho','Uma pessoa responsável + agentes de IA; WIP principal recomendado = 1.'),
 ('Cadência','Sprints por aceite, sem duração fixa. Não converter 72 sprints em 144 semanas.'),
 ('Horizonte','v1.0 completa; v1.1 maturação; v1.2/v2.0 evoluções propostas com ADRs explícitas.'),
 ('Escala do plano',f'{len(SPRINTS)} sprints; 720 tarefas; 72 gates; 19 versões; 37 requisitos; 18 ADRs.'),
 ('Estado inicial','Planejamento. Nenhuma implementação, benchmark ou homologação do app foi realizada nesta entrega.'),
 ('Como começar','Abra Tarefas, filtre Sprint = S01 e siga SL-S01-01. Feche o gate somente com evidência e revisão humana.'),
 ('Colunas editáveis','Status, Responsável, Revisor humano, Data de aceite, Evidência obtida, ADR/justificativa e Observações.'),
 ('Fechamento','OK exige Concluído + revisor + data + evidência + dependências fechadas. PREENCHER CAMPOS NÃO PROVA A EVIDÊNCIA.'),
 ('Dispensas','Somente itens condicionais podem usar Dispensado (ADR), com justificativa/evidência/revisão. Invariantes essenciais não são dispensáveis.'),
 ('Fórmulas','Recalculadas ao abrir em Excel/LibreOffice. Previews podem não calcular fórmulas. Use o status inicial como planejado, não como progresso certificado.'),
 ('Dependências','IDs no nível de tarefa/gate. Ordem numérica é rota solo recomendada; dependências técnicas de sprints estão na aba Sprints.'),
 ('Artefatos','Caminhos listados são saídas futuras no repositório do app. Evidência exigida não significa evidência já produzida.'),
 ('Motores','Recortes upstream, não cópia de produtos completos. Um único backend Qdrant OU HelixDB em produção inicialmente.'),
 ('Offline','Jev/BYOK/web exigem autorização; operação local não depende de chave. Modelos/pacotes devem ser provisionados antes do uso desconectado.'),
 ('Open source','Projeto sem fins comerciais não elimina obrigações upstream ou custos de APIs. Licença open source reconhecida permite uso comercial de terceiros.'),
 ('Prioridade','P0/P1/P2 são relativas ao escopo da versão; um P0 de v2 não antecipa a v2 às dependências da v1.'),
 ('Documento principal',MD.name),
 ('Exportação plana',CSV.name),
 ('Regeneração','Script planejamento/_geracao/render_plan.py preserva campos de acompanhamento por ID. Faça backup; mudanças no plano precisam revalidar dependências.'),
 ('Base','RFC do autor e Sandland-RFC-Integracao-e-Desconstrucao.md, preservados sem alteração.'),
]
for rr,(k,v) in enumerate(info,4):
    readme.cell(rr,1,k).font=Font(name='Calibri',bold=True,color=TEAL,size=11)
    readme.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=8)
    cc=readme.cell(rr,2,v);cc.font=basefont;cc.alignment=Alignment(wrap_text=True,vertical='top')
    readme.row_dimensions[rr].height=38 if len(v)>150 else 30
readme.column_dimensions['A'].width=25
for c in 'BCDEFGH':readme.column_dimensions[c].width=16
readme.freeze_panes='B4'
readme['B20'].hyperlink=MD.name

ws_tasks=tabular('Tarefas',headers,[taskrow(t) for t in all_tasks],[7,19,20,10,33,23,18,49,83,58,91,40,27,22,15,13,51,23,23,18,50,50,47,18,57],'TabelaTarefas')
ws_tasks.freeze_panes='D2'
rowid={t['id']:i+2 for i,t in enumerate(all_tasks)}
for i,t in enumerate(all_tasks,2):
    dep_checks=[f'OR(X{rowid[d]}="OK",X{rowid[d]}="DISPENSADO")' for d in t['dependencies']]
    dep='AND('+','.join(dep_checks)+')' if dep_checks else 'TRUE'
    ws_tasks.cell(i,24,f'=IF(N{i}="Concluído",IF(AND(S{i}<>"",ISNUMBER(T{i}),T{i}>0,V{i}<>"",{dep}),"OK","REVISAR"),IF(N{i}="Dispensado (ADR)",IF(AND(LEFT(Q{i},11)="Condicional",W{i}<>"",S{i}<>"",ISNUMBER(T{i}),T{i}>0,V{i}<>""),"DISPENSADO","REVISAR"),"PENDENTE"))')
    ws_tasks.cell(i,20).number_format='dd/mm/yyyy'
    ws_tasks.row_dimensions[i].height=112 if t['kind']!='Gate' else 98
    for col in [14,18,19,20,22,23,25]:
        ws_tasks.cell(i,col).fill=PatternFill('solid',fgColor='F0FDFA')
    if t['kind']=='Gate':
        for col in [2,7,8]:ws_tasks.cell(i,col).font=Font(name='Calibri',bold=True,color=TEAL,size=11)
statusdv=DataValidation(type='list',formula1='"'+','.join(STATES)+'"',allow_blank=False)
statusdv.errorTitle='Status inválido';statusdv.error='Selecione um status da lista.';statusdv.showErrorMessage=True
ws_tasks.add_data_validation(statusdv);statusdv.add(f'N2:N{len(all_tasks)+1}')
for text,color in [('Concluído',GREEN),('Bloqueado',RED),('Em andamento',AMBER),('Em revisão','DBEAFE')]:
    ws_tasks.conditional_formatting.add(f'N2:N{len(all_tasks)+1}',FormulaRule(formula=[f'N2="{text}"'],fill=PatternFill('solid',fgColor=color)))
ws_tasks.conditional_formatting.add(f'X2:X{len(all_tasks)+1}',FormulaRule(formula=['X2="REVISAR"'],fill=PatternFill('solid',fgColor=RED)))
ws_tasks.conditional_formatting.add(f'X2:X{len(all_tasks)+1}',FormulaRule(formula=['X2="OK"'],fill=PatternFill('solid',fgColor=GREEN)))

last=len(all_tasks)+1
rC=f"Tarefas!$C$2:$C${last}";rD=f"Tarefas!$D$2:$D${last}";rG=f"Tarefas!$G$2:$G${last}";rX=f"Tarefas!$X$2:$X${last}"
release_data=[]
for i,r in enumerate(RELEASES,2):
    ss=[s for s in SPRINTS if s['release']==r[0]]
    nt=sum(len(s['tasks']) for s in ss)
    release_data.append([r[0],r[3],r[4],f'S{r[1]:02d}–S{r[2]:02d}',len(ss),nt,len(ss),r[5],f'=COUNTIFS({rC},A{i},{rG},"<>Gate",{rX},"OK")',f'=COUNTIFS({rC},A{i},{rG},"Gate",{rX},"OK")',f'=COUNTIFS({rC},A{i},{rG},"<>Gate",{rX},"DISPENSADO")',f'=IFERROR(I{i}/(F{i}-K{i}),0)'])
wv=tabular('Versões',['Versão','Nome','Maturidade','Sprints','Nº sprints','Tarefas','Gates','Entregável diferente','Tarefas aceitas','Gates aceitos','Tarefas dispensadas','Progresso aplicável'],release_data,[23,39,33,20,13,12,12,105,18,18,22,22],'TabelaVersoes')
for i in range(2,len(RELEASES)+2):wv.row_dimensions[i].height=68;wv.cell(i,12).number_format='0.0%'
wv.conditional_formatting.add(f'L2:L{len(RELEASES)+1}',DataBarRule(start_type='num',start_value=0,end_type='num',end_value=1,color='0F766E'))

sprint_data=[]
for i,s in enumerate(SPRINTS,2):
    sprint_data.append([s['id'],s['release'],s['title'],s['objective'],s['component'],s['nature'],'; '.join(s['deps']),'; '.join(s['requirements']),s['deliverable'],s['demo'],s['gate_id'],s['risk'],s['note'],f'=COUNTIFS({rD},A{i},{rG},"<>Gate",{rX},"OK")',10,f'=COUNTIFS({rD},A{i},{rG},"<>Gate",{rX},"DISPENSADO")',f'=IFERROR(N{i}/(O{i}-P{i}),0)'])
ws=tabular('Sprints',['Sprint','Versão','Objetivo curto','Objetivo verificável','Componente','Forma','Pré-requisitos técnicos','Requisitos','Entrega','Demonstração de aceite','Gate','Risco','Observação','Tarefas aceitas','Total tarefas','Dispensadas válidas','Progresso aplicável'],sprint_data,[10,22,64,100,40,24,29,33,92,116,22,14,102,19,16,20,19],'TabelaSprints')
for i,s in enumerate(SPRINTS,2):
    ws.row_dimensions[i].height=98;ws.cell(i,17).number_format='0.0%'
    ws.cell(i,11).hyperlink=f"#'Tarefas'!B{rowid[s['gate_id']]}";ws.cell(i,11).font=linkfont

wq=tabular('Requisitos',['ID','Área','Requisito','Origem','Sprints planejados','Verificação'],req_rows,[10,22,111,33,105,125],'TabelaRequisitos')
for i in range(2,len(req_rows)+2):wq.row_dimensions[i].height=86
wa=tabular('ADRs',['ID','Decisão','Sprint','Resultado esperado','Estado','Aprovador','Evidência real'],[list(a)+prior_adrs.get(a[0],['A decidir','Autor','']) for a in ADRS],[17,61,12,123,21,19,63],'TabelaADRs')
for i in range(2,len(ADRS)+2):wa.row_dimensions[i].height=62
we=tabular('Motores',['Motor','Trilha de sprints','Modo de integração','Recorte a reutilizar','Excluir','Desenvolvimento próprio','Referência inicial'],ENGINES,[26,48,67,116,90,99,80],'TabelaMotores')
for i,e in enumerate(ENGINES,2):
    we.row_dimensions[i].height=105
    we.cell(i,7).hyperlink=e[6];we.cell(i,7).font=linkfont
wr=tabular('Riscos',['ID','Risco','Severidade','Onde','Resposta / contingência'],RISKS,[10,116,18,36,143],'TabelaRiscos')
for i in range(2,len(RISKS)+2):wr.row_dimensions[i].height=78

edges=[(d,t['id'],byid[d]['sprint'],t['sprint'],'Bloqueia até aceite; sequência conservadora dentro do sprint') for t in all_tasks for d in t['dependencies']]
wd=tabular('Dependências',['Pré-requisito ID','Item dependente ID','Sprint origem','Sprint destino','Regra'],edges,[24,25,17,17,100],'TabelaDependencias')
for i in range(2,len(edges)+2):wd.row_dimensions[i].height=28

wg=tabular('Gates',['Gate','Sprint','Versão','Demonstração','Entregável','Evidência exigida','Observação'],[[s['gate_id'],s['id'],s['release'],s['demo'],s['deliverable'],f"reports/{s['id']}/acceptance.md",s['note']] for s in SPRINTS],[24,10,22,124,105,50,116],'TabelaGates')
for i in range(2,len(SPRINTS)+2):wg.row_dimensions[i].height=88

# Navegação interna.
rel_row={r[0]:i+2 for i,r in enumerate(RELEASES)}
sprint_row={s['id']:i+2 for i,s in enumerate(SPRINTS)}
for i,t in enumerate(all_tasks,2):
    ws_tasks.cell(i,3).hyperlink=f"#'Versões'!A{rel_row[t['release']]}";ws_tasks.cell(i,3).font=linkfont
    ws_tasks.cell(i,4).hyperlink=f"#'Sprints'!A{sprint_row[t['sprint']]}";ws_tasks.cell(i,4).font=linkfont

panel=wb.create_sheet('Painel',1)
panel.sheet_view.showGridLines=False
panel.merge_cells('A1:H2');panel['A1']='SANDLAND / VISÃO DO ROADMAP';panel['A1'].font=Font(name='Calibri',size=22,bold=True,color=WHITE);panel['A1'].fill=headfill;panel['A1'].alignment=Alignment(vertical='center')
metrics=[('Sprints planejados',72),('Tarefas detalhadas',720),('Gates humanos',72),('Versões/marcos',19),('Requisitos rastreados',37),('ADRs a decidir',18),('Tarefas com fechamento válido',f'=COUNTIFS({rG},"<>Gate",{rX},"OK")'),('Gates com fechamento válido',f'=COUNTIFS({rG},"Gate",{rX},"OK")'),('Itens que exigem revisão',f'=COUNTIF({rX},"REVISAR")')]
for i,(k,v) in enumerate(metrics,4):
    panel.cell(i,1,k).font=Font(name='Calibri',bold=True,color=TEAL,size=12)
    panel.merge_cells(start_row=i,start_column=1,end_row=i,end_column=4)
    panel.cell(i,5,v).font=Font(name='Calibri',bold=True,color=NAVY,size=16)
    panel.row_dimensions[i].height=28
panel.merge_cells('A14:H16');panel['A14']='Estado de entrega: plano, não implementação. Todos os itens começam Não iniciado. Contagem de tarefas não mede horas, semanas ou dificuldade. Fórmulas exigem recálculo no editor de planilhas e não validam a veracidade das evidências.';panel['A14'].alignment=Alignment(wrap_text=True,vertical='center');panel['A14'].font=basefont
for i,name in enumerate(['Tarefas','Sprints','Versões','Motores','Requisitos','ADRs','Gates','Riscos','Dependências'],18):
    panel.cell(i,1,'Abrir '+name);panel.cell(i,1).hyperlink=f"#'{name}'!A1";panel.cell(i,1).font=linkfont
for c in 'ABCDEFGH':panel.column_dimensions[c].width=16
panel.column_dimensions['A'].width=24
chart=BarChart();chart.type='bar';chart.style=10;chart.title='Tarefas planejadas por marco — não duração';chart.y_axis.title='Marco';chart.x_axis.title='Número de tarefas'
chart.add_data(Reference(wv,min_col=6,min_row=1,max_row=len(RELEASES)+1),titles_from_data=True)
chart.set_categories(Reference(wv,min_col=1,min_row=2,max_row=len(RELEASES)+1));chart.height=19;chart.width=27
panel.add_chart(chart,'J4')
wb.active=0
wb.save(XLSX)

# QA dos artefatos de planejamento; não são testes do aplicativo.
check=load_workbook(XLSX,read_only=False,data_only=False)
assert check['Tarefas'].max_row==len(all_tasks)+1
assert len(check['Tarefas'].tables)==1
assert len(check['Sprints'].tables)==1
assert all(check['Tarefas'].cell(rowid[t['id']],24).data_type=='f' for t in all_tasks)
assert all(check['Tarefas'].cell(rowid[t['id']],14).value in STATES for t in all_tasks)
assert check['Tarefas'].freeze_panes=='D2'
check.close()
with CSV.open(encoding='utf-8-sig',newline='') as f:
    csvrows=list(csv.reader(f))
assert len(csvrows)==len(all_tasks)+1 and all(len(r)==len(headers) for r in csvrows)
assert [r[1] for r in csvrows[1:]]==[t['id'] for t in all_tasks]
validation=dict(validated_at='2026-09-22',sprints=len(SPRINTS),implementation_tasks=720,gates=72,total_items=len(all_tasks),releases=len(RELEASES),requirements=len(REQUIREMENTS),adrs=len(ADRS),dependency_edges=len(edges),unique_ids=True,all_dependency_refs_resolve=True,acyclic_graph=True,all_requirements_have_sprints=True,markdown_all_ids_once=True,markdown_fences_balanced=True,csv_rows_match=True,xlsx_rows_and_formulas_checked=True,application_implemented=False,application_tests_executed=False,spreadsheet_formulas_executed_in_excel=False,source_manifest=source_manifest,outputs=[dict(file=p.name,bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in [MD,XLSX,CSV]])
(OUT/'validacao-do-plano.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(validation,ensure_ascii=False,indent=2))
